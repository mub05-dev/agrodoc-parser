# router.py
import os
import platform
import pypdfium2 as pdfium
import pytesseract
from parsers.utils import get_full_text
from collections import defaultdict
from schema import ParseMeta, ParseResponse, ParseResult, BatchSample, AnalysisGroup, BatchResponse
from parsers.registry import PDF_REGISTRY, EXCEL_REGISTRY
import re

_tess = os.getenv("TESSERACT_CMD")
if not _tess and platform.system() == "Windows":
    _win_default = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    if os.path.exists(_win_default):
        _tess = _win_default
if _tess:
    pytesseract.pytesseract.tesseract_cmd = _tess


# ─────────────────────────────────────────────
# Detection
# ─────────────────────────────────────────────

def _detect_from_text(text: str, registry) -> str | None:
    """Itera el registry en orden de prioridad y retorna el laboratory_key del primer match."""
    t = text.upper()
    for parser in registry:
        if parser.detect(t):
            return parser.laboratory_key
    return None


def _get_ocr_text(pdf_path: str) -> str:
    """OCR de la primera página — fallback para PDFs basados en imagen."""
    try:
        pdf = pdfium.PdfDocument(pdf_path)
        try:
            bitmap = pdf[0].render(scale=2)
            image = bitmap.to_pil()
        finally:
            pdf.close()
        return pytesseract.image_to_string(image, lang="eng", config="--psm 6")
    except Exception:
        return ""


def _detect(pdf_path: str) -> tuple[str, bool]:
    """
    Retorna (laboratory_key, used_ocr).
    Intenta detección por texto extraído primero.
    Si el PDF no tiene texto, cae en OCR.
    Lanza ValueError si no se reconoce el laboratorio.
    """
    full_text = get_full_text(pdf_path)
    lab_key = _detect_from_text(full_text, PDF_REGISTRY)

    if lab_key:
        return lab_key, False

    if not full_text.strip():
        ocr_text = _get_ocr_text(pdf_path)
        lab_key = _detect_from_text(ocr_text, PDF_REGISTRY)
        if lab_key:
            return lab_key, True

    raise ValueError(
        f"Laboratory not recognized. "
        f"Supported: {[p.laboratory_key for p in PDF_REGISTRY]}"
    )


# ─────────────────────────────────────────────
# PDF routing
# ─────────────────────────────────────────────

def route_parser(pdf_path: str) -> ParseResponse:
    lab_key, used_ocr = _detect(pdf_path)
    parser = next(p for p in PDF_REGISTRY if p.laboratory_key == lab_key)

    parsed = parser.parse(pdf_path)
    meta = ParseMeta(
        laboratory_key=lab_key,
        parser="ocr" if used_ocr else "deterministic",
        detected_as=lab_key,
    )

    return ParseResponse(
        report=parsed.report,
        samples=parsed.samples,
        total_samples=parsed.total_samples,
        meta=meta,
    )


# ─────────────────────────────────────────────
# Excel routing
# ─────────────────────────────────────────────

def route_excel_parser(xlsx_path: str) -> ParseResponse:
    import openpyxl

    def _detect_excel(path: str) -> str | None:
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet_names = [s.upper().strip() for s in wb.sheetnames]
            wb.close()
            # Construye un texto simulado con los nombres de las hojas
            # para que detect() funcione igual que en PDF
            pseudo_text = " ".join(sheet_names)
            return _detect_from_text(pseudo_text, EXCEL_REGISTRY)
        except Exception:
            return None

    lab_key = _detect_excel(xlsx_path)

    if lab_key is None:
        raise ValueError(
            f"Excel laboratory not recognized. "
            f"Supported: {[p.laboratory_key for p in EXCEL_REGISTRY]}"
        )

    parser = next(p for p in EXCEL_REGISTRY if p.laboratory_key == lab_key)
    parsed = parser.parse(xlsx_path)
    meta = ParseMeta(
        laboratory_key=lab_key,
        parser="deterministic",
        detected_as=lab_key,
    )

    return ParseResponse(
        report=parsed.report,
        samples=parsed.samples,
        total_samples=parsed.total_samples,
        meta=meta,
    )


# ─────────────────────────────────────────────
# Batch
# ─────────────────────────────────────────────

def _sample_sort_key(sample) -> tuple:
    for field in (sample.lab_id, sample.name):
        if field:
            nums = re.findall(r"\d+", str(field))
            if nums:
                return tuple(int(n) for n in nums)
    return (float("inf"),)


def merge_batch_results(results: list) -> BatchResponse:
    laboratory = results[0].report.laboratory
    meta = results[0].meta

    groups: dict = defaultdict(list)
    for result in results:
        atype = result.report.analysis_type or "suelo"
        groups[atype].append(result)

    analyses = []
    grand_total = 0

    for analysis_type, group_results in groups.items():
        documents = []
        indexed_samples = []

        for doc_idx, result in enumerate(group_results):
            doc = {"index": doc_idx, **result.report.model_dump()}
            documents.append(doc)
            for sample in result.samples:
                indexed_samples.append((sample, doc_idx))

        indexed_samples.sort(key=lambda x: _sample_sort_key(x[0]))

        batch_samples = []
        for i, (sample, doc_idx) in enumerate(indexed_samples, 1):
            data = sample.model_dump()
            data["geo_id"] = f"M{i}"
            data["source_document"] = doc_idx
            batch_samples.append(BatchSample(**data))

        analyses.append(AnalysisGroup(
            analysis_type=analysis_type,
            documents=documents,
            samples=batch_samples,
            total_samples=len(batch_samples),
        ))
        grand_total += len(batch_samples)

    return BatchResponse(
        laboratory=laboratory,
        analyses=analyses,
        total_samples=grand_total,
        meta=meta,
    )
