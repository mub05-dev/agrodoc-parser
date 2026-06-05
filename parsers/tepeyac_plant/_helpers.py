"""
tepeyac_plant_parser.py
-----------------------
Parser for Fertilizantes Tepeyac plant (foliar) analysis Excel (.xlsx) reports.
Sheet: "RESULTADOS " — multiple page-blocks, up to 4 samples per block.
"""

import re
import sys
import json
import openpyxl
from datetime import datetime, date
from typing import Optional, List, Tuple


NA_VALUES   = {"na", "n/a", "", "-", "none"}
SAMPLE_COLS = [11, 16, 21, 26]
UNIT_COL    = 6
REF_MIN_COL = 31
REF_MAX_COL = 34

_PARAMS: List[Tuple[str, str, str, str]] = [
    (r"nitr[oó]geno\s+total",    "N",   "%",   "mayores"),
    (r"calcio",                   "Ca",  "%",   "mayores"),
    (r"magnesio",                 "Mg",  "%",   "mayores"),
    (r"potasio",                  "K",   "%",   "mayores"),
    (r"f[oó]sforo|fosforo",      "P",   "%",   "mayores"),
    (r"nitr[aá]tos|n\s*-?\s*no", "NO3", "ppm", "mayores"),
    (r"hierro",                   "Fe",  "ppm", "menores"),
    (r"manganeso",                "Mn",  "ppm", "menores"),
    (r"zinc",                     "Zn",  "ppm", "menores"),
    (r"cobre",                    "Cu",  "ppm", "menores"),
    (r"boro",                     "B",   "ppm", "menores"),
]

_SKIP_LABELS = {
    "elementos mayores", "elementos menores", "unidades",
    "resultados", "resultados ", "referencias",
}


def _cv(val) -> str:
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    return str(val).strip()


def _get(row: tuple, idx: int) -> str:
    return _cv(row[idx]) if idx < len(row) else ""


def _first(row: tuple, *indices: int) -> str:
    for idx in indices:
        v = _get(row, idx)
        if v:
            return v
    return ""


def parse_number(raw) -> Optional[float]:
    s = _cv(raw)
    if s.lower() in NA_VALUES:
        return None
    try:
        return round(float(s.replace(",", ".")), 2)
    except ValueError:
        return None


def get_status(v: Optional[float], rmin: Optional[float], rmax: Optional[float]) -> Optional[str]:
    if v is None:
        return None
    if rmin is not None and v < rmin:
        return "bajo"
    if rmax is not None and v > rmax:
        return "alto"
    if rmin is not None or rmax is not None:
        return "ok"
    return None


def match_param(label: str) -> Tuple[Optional[str], Optional[str], Optional[str]]:
    lc = label.lower().strip()
    if lc in _SKIP_LABELS:
        return None, None, None
    for pat, simbolo, unit, seccion in _PARAMS:
        if re.search(pat, lc):
            return simbolo, unit, seccion
    return None, None, None


def _parse_header(rows: List[tuple]) -> dict:
    """Extract report-level fields from the rows preceding the first sample block."""
    hdr: dict = {}
    for row in rows:
        col0  = _get(row, 0)
        col29 = _get(row, 29)

        if col29 == "Fecha:":
            hdr["fecha_informe"] = _first(row, 32, 33)
        elif "Reporte No" in col29:
            hdr["reporte_no"] = _first(row, 33, 32)
        elif "Solicitud" in col29:
            hdr["solicitud"] = _first(row, 33, 32)

        if col0.startswith("Cliente"):
            hdr["cliente"]          = _get(row, 4)
            hdr["predio"]           = _get(row, 19)
            hdr["cultivo_anterior"] = _first(row, 33, 32)
        elif col0.startswith("Empresa"):
            hdr["empresa"]        = _get(row, 4)
            hdr["riego"]          = _get(row, 19)
            hdr["cultivo_actual"] = _first(row, 33, 32)
        elif col0.lower().startswith("ubicaci"):
            hdr["ubicacion"]       = _get(row, 4)
            hdr["etapa"]           = _get(row, 19)
            hdr["fecha_recepcion"] = _first(row, 32, 33)

    return hdr


def _build_measurement(label: str, simbolo: str, unit: str, seccion: str,
                        value_raw, ref_min_raw, ref_max_raw) -> dict:
    valor = parse_number(value_raw)
    rmin  = parse_number(ref_min_raw)
    rmax  = parse_number(ref_max_raw)
    return {
        "nombre":    label,
        "simbolo":   simbolo,
        "unidad":    unit,
        "valor":     valor,
        "valor_txt": None,
        "rango_min": rmin,
        "rango_max": rmax,
        "status":    get_status(valor, rmin, rmax),
        "seccion":   seccion,
    }


def parse_tepeyac_plant(xlsx_path: str) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    sheet_name = next(
        (s for s in wb.sheetnames if s.strip().upper() == "RESULTADOS"),
        None,
    )
    if sheet_name is None:
        raise ValueError(f"Sheet 'RESULTADOS' not found. Available: {wb.sheetnames}")

    ws = wb[sheet_name]
    all_rows: List[tuple] = [
        tuple(cell.value for cell in row)
        for row in ws.iter_rows()
        if any(cell.value is not None for cell in row)
    ]

    lab_row_indices = [
        i for i, row in enumerate(all_rows)
        if _get(row, 0).strip() == "No. Laboratorio"
    ]

    if not lab_row_indices:
        raise ValueError("No 'No. Laboratorio' rows found in RESULTADOS sheet.")

    header = _parse_header(all_rows[: lab_row_indices[0]])

    all_samples: List[dict] = []

    for block_num, lab_idx in enumerate(lab_row_indices):
        lab_row  = all_rows[lab_idx]
        name_row = all_rows[lab_idx + 1] if lab_idx + 1 < len(all_rows) else ()

        lab_ids      = [_get(lab_row,  c) for c in SAMPLE_COLS]
        sample_names = [_get(name_row, c) for c in SAMPLE_COLS]

        
        active = [i for i, lid in enumerate(lab_ids) if lid and lid.upper() != "NA"]
        if not active:
            continue

        block_samples = [
            {
                "geo_id":           f"M{len(all_samples) + pos + 1}",
                "geo_id_confirmed": False,
                "lab_id":           lab_ids[i],
                "nombre":           sample_names[i] or None,
                "tipo_analisis":    "foliar",
                "mediciones":       [],
            }
            for pos, i in enumerate(active)
        ]

        end_idx  = lab_row_indices[block_num + 1] if block_num + 1 < len(lab_row_indices) else len(all_rows)
        meas_rows = all_rows[lab_idx + 2 : end_idx]

        for row in meas_rows:
            label = _get(row, 0)
            unit_cell = _get(row, UNIT_COL)
            
            if not label or unit_cell.lower() not in ("%", "ppm"):
                continue
            simbolo, unit, seccion = match_param(label)
            if not simbolo:
                continue

            ref_min = _get(row, REF_MIN_COL)
            ref_max = _get(row, REF_MAX_COL)

            for pos, orig_idx in enumerate(active):
                col = SAMPLE_COLS[orig_idx]
                raw = row[col] if col < len(row) else None
                block_samples[pos]["mediciones"].append(
                    _build_measurement(label, simbolo, unit, seccion, raw, ref_min, ref_max)
                )

        all_samples.extend(block_samples)

    return {
        "informe": {
            "laboratorio":   "Fertilizantes Tepeyac",
            "tipo_analisis": "foliar",
            **header,
        },
        "muestras":       all_samples,
        "total_muestras": len(all_samples),
    }
